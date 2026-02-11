import csv
import io
from typing import Optional

from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.core.models import Lead


EXPORT_COLUMNS = [
    "domain", "company_name", "status", "source",
    "is_ecommerce", "ecommerce_platform", "ecommerce_confidence",
    "has_google_analytics", "ga_version", "ga_id",
    "has_gtm", "gtm_id", "has_gsc",
    "has_facebook_pixel", "fb_pixel_id",
    "has_tiktok_pixel", "tiktok_pixel_id",
    "has_hotjar", "has_clarity", "has_mixpanel", "has_segment",
    "hosting_provider", "cdn_provider", "email_provider",
    "is_running_google_ads", "google_ads_count",
    "is_running_meta_ads", "meta_ads_count",
    "is_running_tiktok_ads", "tiktok_ads_count",
    "instagram_url", "instagram_active",
    "facebook_url", "facebook_active",
    "tiktok_url", "tiktok_active",
    "twitter_url", "linkedin_url", "youtube_url",
    "contact_emails", "contact_phones",
    "decision_maker_1_name", "decision_maker_1_title", "decision_maker_1_email",
    "decision_maker_2_name", "decision_maker_2_title", "decision_maker_2_email",
    "gaps_count", "top_gaps",
    "pitch_subject", "pitch_text",
]


def _flatten_lead(lead: Lead) -> dict:
    row = {
        "domain": lead.domain,
        "company_name": lead.company_name or "",
        "status": lead.status,
        "source": lead.source,
    }

    ec = lead.ecommerce
    if ec:
        row["is_ecommerce"] = ec.is_ecommerce
        row["ecommerce_platform"] = ec.platform or ""
        row["ecommerce_confidence"] = ec.confidence_score
    else:
        row["is_ecommerce"] = ""
        row["ecommerce_platform"] = ""
        row["ecommerce_confidence"] = ""

    mt = lead.marketing_tools
    if mt:
        row["has_google_analytics"] = mt.has_google_analytics
        row["ga_version"] = mt.ga_version or ""
        row["ga_id"] = mt.ga_id or ""
        row["has_gtm"] = mt.has_gtm
        row["gtm_id"] = mt.gtm_id or ""
        row["has_gsc"] = mt.has_gsc
        row["has_facebook_pixel"] = mt.has_facebook_pixel
        row["fb_pixel_id"] = mt.fb_pixel_id or ""
        row["has_tiktok_pixel"] = mt.has_tiktok_pixel
        row["tiktok_pixel_id"] = mt.tiktok_pixel_id or ""
        row["has_hotjar"] = mt.has_hotjar
        row["has_clarity"] = mt.has_clarity
        row["has_mixpanel"] = mt.has_mixpanel
        row["has_segment"] = mt.has_segment
    else:
        for col in ["has_google_analytics", "ga_version", "ga_id", "has_gtm", "gtm_id",
                     "has_gsc", "has_facebook_pixel", "fb_pixel_id", "has_tiktok_pixel",
                     "tiktok_pixel_id", "has_hotjar", "has_clarity", "has_mixpanel", "has_segment"]:
            row[col] = ""

    he = lead.hosting_email
    if he:
        row["hosting_provider"] = he.hosting_provider or ""
        row["cdn_provider"] = he.cdn_provider or ""
        row["email_provider"] = he.email_provider or ""
    else:
        row["hosting_provider"] = ""
        row["cdn_provider"] = ""
        row["email_provider"] = ""

    ad = lead.ad_activity
    if ad:
        row["is_running_google_ads"] = ad.is_running_google_ads
        row["google_ads_count"] = ad.google_ads_count
        row["is_running_meta_ads"] = ad.is_running_meta_ads
        row["meta_ads_count"] = ad.meta_ads_count
        row["is_running_tiktok_ads"] = ad.is_running_tiktok_ads
        row["tiktok_ads_count"] = ad.tiktok_ads_count
    else:
        for col in ["is_running_google_ads", "google_ads_count", "is_running_meta_ads",
                     "meta_ads_count", "is_running_tiktok_ads", "tiktok_ads_count"]:
            row[col] = ""

    sm = lead.social_media
    if sm:
        row["instagram_url"] = sm.instagram_url or ""
        row["instagram_active"] = sm.instagram_active if sm.instagram_active is not None else ""
        row["facebook_url"] = sm.facebook_url or ""
        row["facebook_active"] = sm.facebook_active if sm.facebook_active is not None else ""
        row["tiktok_url"] = sm.tiktok_url or ""
        row["tiktok_active"] = sm.tiktok_active if sm.tiktok_active is not None else ""
        row["twitter_url"] = sm.twitter_url or ""
        row["linkedin_url"] = sm.linkedin_url or ""
        row["youtube_url"] = sm.youtube_url or ""
    else:
        for col in ["instagram_url", "instagram_active", "facebook_url", "facebook_active",
                     "tiktok_url", "tiktok_active", "twitter_url", "linkedin_url", "youtube_url"]:
            row[col] = ""

    # Contacts
    emails = [c.value for c in lead.contacts if c.type == "email"]
    phones = [c.value for c in lead.contacts if c.type == "phone"]
    row["contact_emails"] = "; ".join(emails)
    row["contact_phones"] = "; ".join(phones)

    # Decision makers
    for i, dm in enumerate(lead.decision_makers[:2]):
        idx = i + 1
        row[f"decision_maker_{idx}_name"] = dm.name or ""
        row[f"decision_maker_{idx}_title"] = dm.title or ""
        row[f"decision_maker_{idx}_email"] = dm.email or ""
    for i in range(len(lead.decision_makers[:2]), 2):
        idx = i + 1
        row[f"decision_maker_{idx}_name"] = ""
        row[f"decision_maker_{idx}_title"] = ""
        row[f"decision_maker_{idx}_email"] = ""

    # Pitch
    p = lead.pitch
    if p:
        row["gaps_count"] = len(p.gaps_identified) if p.gaps_identified else 0
        row["top_gaps"] = "; ".join(p.gaps_identified[:5]) if p.gaps_identified else ""
        row["pitch_subject"] = p.pitch_subject_line or ""
        row["pitch_text"] = p.pitch_text or ""
    else:
        row["gaps_count"] = ""
        row["top_gaps"] = ""
        row["pitch_subject"] = ""
        row["pitch_text"] = ""

    return row


async def _get_leads(db: AsyncSession, lead_ids: Optional[list[int]] = None) -> list[Lead]:
    query = (
        select(Lead)
        .options(
            selectinload(Lead.ecommerce),
            selectinload(Lead.marketing_tools),
            selectinload(Lead.hosting_email),
            selectinload(Lead.ad_activity),
            selectinload(Lead.social_media),
            selectinload(Lead.decision_makers),
            selectinload(Lead.contacts),
            selectinload(Lead.pitch),
        )
    )
    if lead_ids:
        query = query.where(Lead.id.in_(lead_ids))
    result = await db.execute(query.order_by(Lead.id))
    return list(result.scalars().all())


async def export_csv(db: AsyncSession, lead_ids: Optional[list[int]] = None) -> str:
    leads = await _get_leads(db, lead_ids)
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=EXPORT_COLUMNS)
    writer.writeheader()
    for lead in leads:
        writer.writerow(_flatten_lead(lead))
    return output.getvalue()


async def export_excel(db: AsyncSession, lead_ids: Optional[list[int]] = None) -> bytes:
    leads = await _get_leads(db, lead_ids)
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"

    # Header row
    for col_idx, header in enumerate(EXPORT_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = cell.font.copy(bold=True)

    # Data rows
    for row_idx, lead in enumerate(leads, 2):
        data = _flatten_lead(lead)
        for col_idx, col_name in enumerate(EXPORT_COLUMNS, 1):
            ws.cell(row=row_idx, column=col_idx, value=data.get(col_name, ""))

    # Auto-width columns
    for col_idx, header in enumerate(EXPORT_COLUMNS, 1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max(
            len(header) + 2, 15
        )

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
